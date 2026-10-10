"""
recheck_cq8.py -- after the route-1 tightening: (a) which audited CQ8 pairs survive in
the rebuilt graphs, cross-tabulated with their audit verdict, and (b) a FRESH random
sample of CQ8 pairs from the rebuilt graphs for a new audit. No LLM calls.

The original audit evaluates the original graph; it cannot establish the revised
graph's accuracy -- hence the fresh sample (b).

Run from the project root, after build_kg.py (rebuilt) and with your CHECKED audit file:
    python3 -m src.kg.recheck_cq8 data/interim/kg/audit_sheets_checked.xlsx [N]

Outputs (data/interim/kg/):
    cq8_recheck.csv            every audited pair: config, verdict, retained yes/no
    audit_cq8_v2.xlsx          fresh sample, N per configuration (default 25), same
                               columns as the original cq8_pairs sheet
    printed: verdict x retained table and CQ8 counts before/after
"""

import csv
import json
import os
import pickle
import random
import sys
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src.kg.build_kg import cq8_pairs

D = "data/interim"
KG = f"{D}/kg"
CONFIGS = ["qwen_single", "qwen_iter", "llama_single", "llama_iter"]


def main():
    audit_path = sys.argv[1] if len(sys.argv) > 1 else f"{KG}/audit_sheets_checked.xlsx"
    n = int(sys.argv[2]) if len(sys.argv) > 2 else 25
    rnd = random.Random(7)

    rows = list(load_workbook(audit_path, read_only=True)["cq8_pairs"].values)
    H = list(rows[0])
    audited = [dict(zip(H, r)) for r in rows[1:] if r[0]]

    pairs, graphs = {}, {}
    for c in CONFIGS:
        with open(f"{KG}/kg_{c}_full.gpickle", "rb") as f:
            G = pickle.load(f)
        graphs[c] = G
        pairs[c] = cq8_pairs(G)

    # (a) audited pairs: retained?
    out, tab = [], Counter()
    for a in audited:
        c, d1, d2 = a["config"], a["domain_1"], a["domain_2"]
        kept = any((d1 in x[2] and d2 in y[2]) or (d2 in x[2] and d1 in y[2])
                   for x, y, _ in pairs[c])
        verdict = (a.get("verdict") or "").strip().lower() or "(empty)"
        tab[(verdict, kept)] += 1
        out.append({"audit_id": a["audit_id"], "config": c, "verdict": verdict,
                    "retained": "yes" if kept else "no", "product": a["product"]})
    with open(f"{KG}/cq8_recheck.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(out[0].keys()))
        w.writeheader()
        w.writerows(out)
    print("AUDITED CQ8 PAIRS IN THE REBUILT GRAPHS (verdict x retained)")
    for v in sorted({k[0] for k in tab}):
        print(f"  {v:<10} retained {tab[(v, True)]:>3} | removed {tab[(v, False)]:>3}")

    # CQ8 counts before / after
    before = {}
    if os.path.exists(f"{KG}/summary_v1.csv"):
        with open(f"{KG}/summary_v1.csv", encoding="utf-8") as f:
            for r in csv.DictReader(f):
                if r["page_set"] == "full":
                    before[r["config"]] = r.get("cq8_candidate_pairs")
    print("\nCQ8 CANDIDATE PAIRS before -> after")
    for c in CONFIGS:
        print(f"  {c:<13} {before.get(c, '?')} -> {len(pairs[c])}")

    # (b) fresh sample of retained pairs, excluding already audited ones
    done = {(a["config"], frozenset((a["domain_1"], a["domain_2"]))) for a in audited}
    sample, want = [], set()
    nm = lambda G, x: G.nodes[x].get("name", x)
    for c in CONFIGS:
        G = graphs[c]
        pool = [p for p in pairs[c] if (c, frozenset((p[0][2][0], p[1][2][0]))) not in done]
        for (c1, r1, d1), (c2, r2, d2), p in rnd.sample(pool, min(n, len(pool))):
            want.update([d1[0], d2[0]])
            sample.append({"config": c, "product": nm(G, p),
                           "product_identity": G.nodes[p].get("identity"),
                           "company_1": f"{nm(G, c1)} ({r1})", "domain_1": d1[0],
                           "company_2": f"{nm(G, c2)} ({r2})", "domain_2": d2[0]})
    text = {}
    with open(f"{D}/fetched_pages.jsonl", encoding="utf-8") as f:
        for line in f:
            try:
                r = json.loads(line)
            except json.JSONDecodeError:
                continue
            if r.get("domain") in want and r["domain"] not in text:
                text[r["domain"]] = (r.get("text") or "")[:30000]
    wb = Workbook()
    ws = wb.active
    ws.title = "cq8_pairs_v2"
    cols = ["audit_id"] + list(sample[0].keys()) + ["page_text_1", "page_text_2", "verdict", "note"]
    ws.append(cols)
    for i, s in enumerate(sample, 1):
        ws.append([f"cq8v2_{i:03d}"] + list(s.values()) +
                  [text.get(s["domain_1"], ""), text.get(s["domain_2"], ""), "", ""])
    yellow = PatternFill("solid", start_color="FFF2CC")
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row > 1 and cell.column >= len(cols) - 1:
                cell.fill = yellow
    dv = DataValidation(type="list", formula1='"correct,incorrect,unclear"', allow_blank=True)
    ws.add_data_validation(dv)
    col = ws.cell(row=1, column=len(cols) - 1).column_letter
    dv.add(f"{col}2:{col}{len(sample) + 1}")
    wb.save(f"{KG}/audit_cq8_v2.xlsx")
    print(f"\n[DONE] {KG}/cq8_recheck.csv and {KG}/audit_cq8_v2.xlsx ({len(sample)} fresh pairs)")


if __name__ == "__main__":
    main()
