"""
inter_annotator_agreement.py

Computes inter-annotator agreement (percent agreement + Cohen's Kappa)
between Claude's annotation and the user's independent gold annotation.
"""

import sys
import openpyxl
from collections import Counter


def normalize(s):
    if not s:
        return "unknown"
    s = str(s).strip().lower()
    return s if s else "unknown"


def binarize(val, field):
    if field == "manufactures":
        return val
    return "unknown" if val == "unknown" else "named"


def cohens_kappa(pairs):
    n = len(pairs)
    if n == 0:
        return None, None
    labels = sorted(set([p[0] for p in pairs] + [p[1] for p in pairs]))

    agree = sum(1 for a, b in pairs if a == b)
    po = agree / n

    r1_counts = Counter(p[0] for p in pairs)
    r2_counts = Counter(p[1] for p in pairs)
    pe = sum((r1_counts[l] / n) * (r2_counts[l] / n) for l in labels)

    if pe == 1.0:
        kappa = 1.0 if po == 1.0 else 0.0
    else:
        kappa = (po - pe) / (1 - pe)

    return po, kappa


def interpret_kappa(k):
    if k is None:
        return "N/A"
    if k < 0:
        return "Poor (worse than chance)"
    elif k < 0.20:
        return "Slight"
    elif k < 0.40:
        return "Fair"
    elif k < 0.60:
        return "Moderate"
    elif k < 0.80:
        return "Substantial"
    else:
        return "Almost Perfect"


def main():
    if len(sys.argv) != 2:
        print("Usage: python3 inter_annotator_agreement.py <Annotation118.xlsx>")
        sys.exit(1)

    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]

    header = [cell.value for cell in ws[1]]
    col_idx = {name: i for i, name in enumerate(header)}

    fields = {
        "manufactures": ("claude_manufactures_yes_no", "GOLD_claude_manufactures_yes_no"),
        "distributed_by": ("claude_distributed_by", "GOLD_claude_distributed_by"),
        "made_in": ("claude_made_in", "GOLD_claude_made_in"),
        "is_component_of": ("claude_is_component_of", "GOLD_claude_is_component_of"),
    }

    all_pairs = {f: [] for f in fields}
    n_rows = 0

    for row_idx in range(2, ws.max_row + 1):
        rid = ws.cell(row=row_idx, column=col_idx["record_id"] + 1).value
        if not rid:
            continue
        n_rows += 1
        for field, (claude_col, gold_col) in fields.items():
            claude_val = normalize(ws.cell(row=row_idx, column=col_idx[claude_col] + 1).value)
            gold_val = normalize(ws.cell(row=row_idx, column=col_idx[gold_col] + 1).value)
            all_pairs[field].append((claude_val, gold_val))

    print(f"[INFO] Loaded {n_rows} annotated rows\n")
    print("="*90)
    print("INTER-ANNOTATOR AGREEMENT: Claude vs. User's Independent Gold Annotation")
    print("="*90)

    for field, pairs in all_pairs.items():
        po_exact, kappa_exact = cohens_kappa(pairs)

        bin_pairs = [(binarize(a, field), binarize(b, field)) for a, b in pairs]
        po_bin, kappa_bin = cohens_kappa(bin_pairs)

        print(f"\n-- {field} --")
        print(f"  Exact match:      {po_exact:.3f} agreement, Cohen's Kappa = {kappa_exact:.3f} ({interpret_kappa(kappa_exact)})")
        if field != "manufactures":
            print(f"  Binarized (unknown vs. named): {po_bin:.3f} agreement, Cohen's Kappa = {kappa_bin:.3f} ({interpret_kappa(kappa_bin)})")

        disagreements = [(a, b) for a, b in pairs if a != b]
        print(f"  Disagreements: {len(disagreements)} / {len(pairs)}")
        if disagreements:
            print(f"  Sample disagreements (claude, gold): {disagreements[:5]}")


if __name__ == "__main__":
    main()
