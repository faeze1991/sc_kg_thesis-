"""
convert_gold_xlsx.py

Converts the user's own independent gold annotation xlsx (containing
GOLD_* columns) into a CSV formatted for build_gold118_scoring.py, using
the GOLD_* columns as the true gold standard.

Usage:
    python3 convert_gold_xlsx.py <input.xlsx> <output.csv>
"""

import sys
import csv
import openpyxl


def main():
    if len(sys.argv) != 3:
        print("Usage: python3 convert_gold_xlsx.py <input.xlsx> <output.csv>")
        sys.exit(1)

    input_path, output_path = sys.argv[1], sys.argv[2]

    wb = openpyxl.load_workbook(input_path, data_only=True)
    ws = wb["Sheet1"]

    header = [cell.value for cell in ws[1]]
    col_idx = {name: i for i, name in enumerate(header)}

    rows_out = []
    n_filled_blank = 0
    for row_idx in range(2, ws.max_row + 1):
        row_vals = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]
        rid = row_vals[col_idx["record_id"]]
        if not rid:
            continue

        gold_manuf = row_vals[col_idx["GOLD_claude_manufactures_yes_no"]]
        gold_dist = row_vals[col_idx["GOLD_claude_distributed_by"]]
        gold_made = row_vals[col_idx["GOLD_claude_made_in"]]
        gold_comp = row_vals[col_idx["GOLD_claude_is_component_of"]]

        filled = []
        if not gold_manuf:
            gold_manuf, filled = "unknown", filled + ["manufactures"]
        if not gold_dist:
            gold_dist, filled = "unknown", filled + ["distributed_by"]
        if not gold_made:
            gold_made, filled = "unknown", filled + ["made_in"]
        if not gold_comp:
            gold_comp, filled = "unknown", filled + ["is_component_of"]
        if filled:
            n_filled_blank += 1
            print(f"[WARN] {rid}: filled blank field(s) with 'unknown': {filled}")

        rows_out.append({
            "record_id": rid,
            "gold_manufactures_yes_no": str(gold_manuf).strip().lower(),
            "gold_distributed_by": str(gold_dist).strip().lower(),
            "gold_made_in": str(gold_made).strip().lower(),
            "gold_is_component_of": str(gold_comp).strip().lower(),
        })

    fieldnames = ["record_id", "gold_manufactures_yes_no", "gold_distributed_by",
                  "gold_made_in", "gold_is_component_of"]
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows_out)

    print(f"\n[DONE] Wrote {len(rows_out)} rows to {output_path}")
    print(f"[INFO] {n_filled_blank} row(s) had blank gold cells filled with 'unknown'")


if __name__ == "__main__":
    main()
