"""
classical_baselines.py -- classical ML baselines (TF-IDF + SVM / Logistic Regression)
for supply chain relation extraction, following the classical-baseline paradigm of
Kim et al. (2025) (SVM over bag-of-words features).

What changed compared with the earlier version
  * Test set = the 320 stratified gold pages (data/interim/gold_annotation.xlsx).
  * The script no longer computes its own metric (macro-F1). It writes predictions
    in the same format as the LLM outputs, and score_unified.py scores them with
    exactly the same metric as the LLMs:
        python3 -m src.kg.score_unified \
            --extra classical_svm=data/interim/eval/preds_classical_svm.jsonl \
            --extra classical_lr=data/interim/eval/preds_classical_lr.jsonl
  * Both training and test text are cut to the first 3,000 characters, the same
    input the LLMs received.

Scope (unchanged)
  * manufactures: 3-class classification (yes / no / unknown).
  * has_brand (distributed_by), made_in, is_component_of: binary detection only
    (unknown vs. a value is stated). A fixed-label classifier cannot output a
    company or place name, so for these fields the placeholder "__named__" is
    written and ONLY the "binarized" scores are meaningful for the baselines.
  * company / product are not predicted (left out -> scored as unknown).

Training data ("silver" labels)
  Claude's annotations in annotation_sample_WITH_CLAUDE.csv (366 rows) and
  annotation_topup_WITH_CLAUDE.csv (800 rows), columns claude_*. These were made
  under an earlier, looser version of the guidelines; the gold set follows the
  final strict guidelines (documented limitation). Any training page whose domain
  is in the gold set is removed (no leakage); the script prints how many.

Usage (from the project root):
    python3 -m src.kg.classical_baselines
"""

import csv
import json
import os

from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.svm import LinearSVC

D = "data/interim"
TRAIN_FILES = [f"{D}/annotation_sample_WITH_CLAUDE.csv", f"{D}/annotation_topup_WITH_CLAUDE.csv"]
GOLD = f"{D}/gold_annotation.xlsx"
GOLD_TEXT = f"{D}/annotation_strat_BLIND.xlsx"  # page_text of the 320 gold pages
OUT_DIR = f"{D}/eval"
TEXT_LIMIT = 3000  # same input length as the LLMs

# field in the output (= key in the LLM files) -> silver label column
FIELDS = {
    "manufactures": "claude_manufactures_yes_no",
    "distributed_by": "claude_distributed_by",   # = has_brand
    "made_in": "claude_made_in",
    "is_component_of": "claude_is_component_of",
}
NEG = {"", "unknown", "none", "n/a", "na"}


def manu_label(v):
    v = str(v or "").strip().lower()
    if v.startswith("yes"):
        return "yes"
    if v.startswith("no"):
        return "no"
    return "unknown"


def bin_label(v):
    return "unknown" if str(v or "").strip().lower() in NEG else "__named__"


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        sample = f.read(4096)
        f.seek(0)
        delim = ";" if sample.count(";") > sample.count(",") else ","
        return list(csv.DictReader(f, delimiter=delim))


def main():
    # gold pages: domains + text
    gold_rows = list(load_workbook(GOLD, read_only=True).worksheets[0].values)
    gh = [str(h) for h in gold_rows[0]]
    gold = [dict(zip(gh, r)) for r in gold_rows[1:] if r[0]]
    gold_domains = {str(g["domain"]).strip() for g in gold}
    text_rows = list(load_workbook(GOLD_TEXT, read_only=True)["Annotation"].values)
    th = [str(h) for h in text_rows[0]]
    gtext = {r[th.index("record_id")]: (r[th.index("page_text")] or "") for r in text_rows[1:]}
    test = [(g["record_id"], str(g["domain"]).strip(), gtext.get(g["record_id"], "")[:TEXT_LIMIT])
            for g in gold]
    print(f"[INFO] test: {len(test)} gold pages ({sum(1 for t in test if t[2])} with text)")

    # silver training data
    train, seen, n_leak = [], set(), 0
    for path in TRAIN_FILES:
        for r in read_csv(path):
            dom = (r.get("domain") or "").strip()
            text = r.get("archived_page_text") or ""
            if not dom or not text or dom in seen:
                continue
            if dom in gold_domains:
                n_leak += 1
                continue
            seen.add(dom)
            train.append((text[:TEXT_LIMIT], {f: r.get(c, "") for f, c in FIELDS.items()}))
    print(f"[INFO] training: {len(train)} silver-labelled pages "
          f"({n_leak} removed because their domain is in the gold set)")

    vec = TfidfVectorizer(max_features=5000, ngram_range=(1, 2), stop_words="english")
    X_train = vec.fit_transform([t for t, _ in train])
    X_test = vec.transform([t for _, _, t in test])

    classifiers = {
        "classical_svm": lambda: LinearSVC(class_weight="balanced", max_iter=5000),
        "classical_lr": lambda: LogisticRegression(max_iter=1000, class_weight="balanced"),
    }
    os.makedirs(OUT_DIR, exist_ok=True)
    for name, make in classifiers.items():
        preds = {rid: {"domain": dom} for rid, dom, _ in test}
        for field in FIELDS:
            to_label = manu_label if field == "manufactures" else bin_label
            y = [to_label(lab[field]) for _, lab in train]
            classes = sorted(set(y))
            if len(classes) < 2:
                print(f"[WARN] {name} {field}: only one class in training data; predicting it")
                pred = [classes[0]] * len(test)
            else:
                clf = make()
                clf.fit(X_train, y)
                pred = [str(p) for p in clf.predict(X_test)]
            dist = {c: y.count(c) for c in classes}
            pdist = {c: pred.count(c) for c in sorted(set(pred))}
            print(f"[INFO] {name:<14} {field:<16} train {dist}  ->  test predictions {pdist}")
            for (rid, _, _), p in zip(test, pred):
                preds[rid][field] = p
        out = f"{OUT_DIR}/preds_{name}.jsonl"
        with open(out, "w", encoding="utf-8") as f:
            for rec in preds.values():
                f.write(json.dumps(rec, ensure_ascii=False) + "\n")
        print(f"[DONE] {out}")

    print("\nNext: python3 -m src.kg.score_unified "
          f"--extra classical_svm={OUT_DIR}/preds_classical_svm.jsonl "
          f"--extra classical_lr={OUT_DIR}/preds_classical_lr.jsonl")
    print("For the baselines, read only the BINARIZED scores of has_brand, made_in and "
          "is_component_of (they cannot name a value); manufactures is comparable in all modes.")


if __name__ == "__main__":
    main()
