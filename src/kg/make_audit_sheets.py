"""
make_audit_sheets.py -- sampled graph-level audits (step 5c). No LLM calls.

Run from the project root (after build_kg.py):
    python3 -m src.kg.make_audit_sheets          # 25 per configuration (default)
    python3 -m src.kg.make_audit_sheets 15       # smaller samples

Writes data/interim/kg/audit_sheets.xlsx with four sheets. For every row, check
the statement against the page text and fill `verdict` (correct / incorrect /
unclear) and, if useful, `note`.
  component_edges  random is_component_of edges, N per configuration (full set)
  cq8_pairs        random CQ8 candidate pairs, N per configuration: are both links
                   valid AND is it the same product? Only then may a pair be
                   called a "hidden" connection.
  id_merges        30 random accepted identifier (WDC cluster) merges: same product?
  offers           20 random offers edges (page-context-derived): does the page
                   really offer this product, and is the operator right?
Seed is fixed (42), so the sample is reproducible.
"""

import csv
import json
import pickle
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src.kg.build_kg import cq8_pairs

D = "data/interim"
KG = f"{D}/kg"
CONFIGS = ["qwen_single", "qwen_iter", "llama_single", "llama_iter"]
TEXT_CAP = 30000


def main():
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 25
    rnd = random.Random(42)
    sheets = {"component_edges": [], "cq8_pairs": [], "id_merges": [], "offers": []}
    want = set()

    graphs = {}
    for c in CONFIGS:
        try:
            with open(f"{KG}/kg_{c}_full.gpickle", "rb") as f:
                graphs[c] = pickle.load(f)
        except FileNotFoundError:
            print(f"[WARN] kg_{c}_full.gpickle missing, skipped")
    nm = lambda G, x: G.nodes[x].get("name", x)

    for c, G in graphs.items():
        comp = [(u, v, d) for u, v, d in G.edges(data=True) if d["relation"] == "is_component_of"]
        for u, v, d in rnd.sample(comp, min(n, len(comp))):
            dom = d["domains"][0]
            want.add(dom)
            sheets["component_edges"].append({
                "config": c, "statement": f"{nm(G, u)}  --is_component_of-->  {nm(G, v)}",
                "parent_identity": G.nodes[v].get("identity"), "source_count": d["source_count"],
                "domain": dom, "url": (d["sample_urls"] or [""])[0], "_text": dom})
        pairs = cq8_pairs(G)
        for (c1, r1, d1), (c2, r2, d2), p in rnd.sample(pairs, min(n, len(pairs))):
            want.update([d1[0], d2[0]])
            sheets["cq8_pairs"].append({
                "config": c, "product": nm(G, p), "product_identity": G.nodes[p].get("identity"),
                "company_1": f"{nm(G, c1)} ({r1})", "domain_1": d1[0],
                "company_2": f"{nm(G, c2)} ({r2})", "domain_2": d2[0],
                "_text": d1[0], "_text2": d2[0]})

    # identifier merges (config-independent source: WDC clusters) from qwen_single
    try:
        with open(f"{KG}/entity_matches_qwen_single_full.csv", encoding="utf-8") as f:
            acc = [r for r in csv.DictReader(f) if r["decision"] == "accepted" and r["route"] == "1_identifier"]
        titles = {}
        with open(f"{D}/wdc_fields.jsonl", encoding="utf-8") as f:
            for line in f:
                r = json.loads(line)
                titles[r["domain"]] = r.get("wdc_title") or ""
        for r in rnd.sample(acc, min(30, len(acc))):
            doms = [x.strip() for x in r["domains"].split("|")]
            sheets["id_merges"].append({
                "cluster": r["key"], "n_pages": r["n_pages"], "brands": r["brands"],
                "extracted_names": r["names"],
                "wdc_titles": " || ".join(f"{d}: {titles.get(d, '')}" for d in doms[:6])})
    except FileNotFoundError:
        print("[WARN] entity_matches_qwen_single_full.csv missing: id_merges sheet empty")

    if "qwen_single" in graphs:
        G = graphs["qwen_single"]
        off = [(u, v, d) for u, v, d in G.edges(data=True) if d["relation"] == "offers"]
        for u, v, d in rnd.sample(off, min(20, len(off))):
            dom = d["domains"][0]
            want.add(dom)
            sheets["offers"].append({"statement": f"{nm(G, u)}  --offers-->  {nm(G, v)}",
                                     "domain": dom, "url": (d["sample_urls"] or [""])[0], "_text": dom})

    text = {}
    with open(f"{D}/fetched_pages.jsonl", encoding="utf-8") as f:
        for line in f:
            try:
                r = json.loads(line)
            except json.JSONDecodeError:
                continue
            if r.get("domain") in want and r["domain"] not in text:
                text[r["domain"]] = (r.get("text") or "")[:TEXT_CAP]

    wb = Workbook()
    wb.remove(wb.active)
    yellow = PatternFill("solid", start_color="FFF2CC")
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        if not rows:
            ws.append(["(no rows)"])
            continue
        cols = [k for k in rows[0] if not k.startswith("_")]
        has_t2 = "_text2" in rows[0]
        if "_text" in rows[0]:
            cols.append("page_text" if not has_t2 else "page_text_1")
        if has_t2:
            cols.append("page_text_2")
        cols += ["verdict", "note"]
        ws.append(["audit_id"] + cols)
        for i, r in enumerate(rows, 1):
            vals = [r[k] for k in r if not k.startswith("_")]
            if "_text" in r:
                vals.append(text.get(r["_text"], ""))
            if "_text2" in r:
                vals.append(text.get(r["_text2"], ""))
            ws.append([f"{name}_{i:03d}"] + vals + ["", ""])
        last = len(cols) + 1
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.row > 1 and cell.column >= last - 1:
                    cell.fill = yellow
        for j in range(1, last + 1):
            ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = \
                70 if "text" in str(ws.cell(row=1, column=j).value) else 26
        dv = DataValidation(type="list", formula1='"correct,incorrect,unclear"', allow_blank=True)
        ws.add_data_validation(dv)
        col = ws.cell(row=1, column=last - 1).column_letter
        dv.add(f"{col}2:{col}{len(rows) + 1}")
        ws.freeze_panes = "B2"
        print(f"[INFO] {name}: {len(rows)} rows")
    wb.save(f"{KG}/audit_sheets.xlsx")
    print(f"[DONE] {KG}/audit_sheets.xlsx")


if __name__ == "__main__":
    main()
